import sys,unittest,tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from validator import pass_rules,DEFAULTS,analyze
from openpyxl import Workbook
class Tests(unittest.TestCase):
 def test_valid_key(self):self.assertFalse(pass_rules('K7m!R4t%H5x+',DEFAULTS))
 def test_banned(self):self.assertIn('KEY_BANNED_CHARACTER',pass_rules('K7m!R4t%H5x0',DEFAULTS))
 def test_length(self):self.assertIn('KEY_LENGTH',pass_rules('Ab3!',DEFAULTS))
 def test_missing_upper(self):self.assertIn('KEY_NO_UPPERCASE',pass_rules('aa3!aa3!aa3!',DEFAULTS))
 def test_pipeline(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'t.xlsx';w=Workbook();s=w.active;s.append(['CT S/N','MAC','WiFi 2.4G SSID','WiFi 5G SSID','WiFi 2.4G Key','WiFi 5G Key']);s.append(['26A4372UXXF-CF000001','A01842F8D4B7','DIGIFIBRA-Ab3C','DIGIFIBRA-PLUS-Ab3C','K7m!R4t%H5x+','K7m!R4t%H5x+']);s.append(['26A4372UXXF-CF000001','A01842F8D4B7','DIGIFIBRA-Ab3C','DIGIFIBRA-PLUS-Ab3C','K7m!R4t%H5x+','K7m!R4t%H5x+']);w.save(p)
   result=analyze(p,d,{'check_gpon_relation':False});self.assertEqual(result['total'],2);self.assertEqual(result['fail'],2);self.assertIn('KEY_DUPLICATE',result['errors']);self.assertTrue(Path(result['report']).exists())
 def test_ssid_dup_allowed_and_mac_mapping(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'t.xlsx';w=Workbook();s=w.active;s.append(['CT S/N','MAC','WiFi 2.4G SSID','WiFi 5G SSID','WiFi 2.4G Key','WiFi 5G Key'])
   for sn,mac,key in [(1,'A01842F8D4B7','K7m!R4t%H5x+'),(3,'A01842F8D4CB','R4t%H5x+K7m!')]:
    s.append([f'26A4372UXXF-CF{sn:06d}',mac,'DIGIFIBRA-Ab3C','DIGIFIBRA-PLUS-Ab3C',key,key])
   w.save(p);r=analyze(p,d,{'check_gpon_relation':False});self.assertEqual(r['fail'],0,r['errors']);self.assertNotIn('SSID_DUPLICATE',r['errors'])
 def test_mapping_and_block_overlap(self):
  with tempfile.TemporaryDirectory() as d:
   p=Path(d)/'t.xlsx';w=Workbook();s=w.active;s.append(['CT S/N','MAC','WiFi 2.4G SSID','WiFi 5G SSID','WiFi 2.4G Key','WiFi 5G Key'])
   for sn,mac,key in [(1,'A01842F8D4B7','K7m!R4t%H5x+'),(2,'A01842F8D4BE','R4t%H5x+K7m!')]:s.append([f'26A4372UXXF-CF{sn:06d}',mac,'DIGIFIBRA-Ab3C','DIGIFIBRA-PLUS-Ab3C',key,key])
   w.save(p);r=analyze(p,d,{'check_gpon_relation':False});self.assertEqual(r['fail'],2,r['errors']);self.assertIn('MAC_BLOCK_OVERLAP',r['errors']);self.assertIn('MAC_SN_MAPPING',r['errors']);self.assertIn('MAC_STEP_ALIGNMENT',r['errors'])
if __name__=='__main__':unittest.main()

